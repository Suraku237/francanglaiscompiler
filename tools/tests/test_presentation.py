"""Regression checks for the concise and reference editable presentations."""

import hashlib
import json
from pathlib import Path
import unittest
from zipfile import ZipFile

import pymupdf
from pptx import Presentation
from pptx.enum.action import PP_ACTION
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame

from tools.build_presentation import (
    BODY_FONT, BRIEF, DIRECTORY, OUTPUT, ROOT, compact, fitted_lines, load_specs, sha256,
    validate_pptx, wrap,
)


class PresentationTests(unittest.TestCase):
    def test_exactly_one_hundred_source_grounded_slides(self) -> None:
        specs = load_specs()
        self.assertEqual(len(specs), 100)
        self.assertEqual(len({spec.title for spec in specs}), 100)
        self.assertTrue(all(spec.notes and spec.sources for spec in specs))

    def test_all_file_map_cards_point_to_real_files(self) -> None:
        for spec in load_specs():
            if spec.layout == "files":
                for source, _ in spec.points:
                    with self.subTest(source=source):
                        self.assertTrue((ROOT / source).is_file())

    def test_actual_powerpoint_is_editable_and_has_notes(self) -> None:
        result = validate_pptx(OUTPUT)
        self.assertEqual(result["slides"], 100)
        self.assertEqual(result["speaker_note_pages"], 100)
        self.assertGreater(result["editable_text_shapes"], 1000)
        self.assertGreaterEqual(result["picture_placements"], 15)
        self.assertEqual(result["off_canvas_shapes"], 0)

    def test_published_titles_and_notes_match_the_content(self) -> None:
        presentation = Presentation(str(OUTPUT))
        for number, (spec, slide) in enumerate(zip(load_specs(), presentation.slides), 1):
            with self.subTest(slide=number):
                titles = [
                    shape.text for shape in slide.shapes
                    if isinstance(shape, Shape) and shape.name.startswith("title-")
                ]
                self.assertEqual(len(titles), 1)
                self.assertEqual(compact(titles[0]), compact(spec.title))
                notes = slide.notes_slide.notes_text_frame
                self.assertIsNotNone(notes)
                if notes is not None:
                    self.assertIn(spec.notes, notes.text)
                    for source in spec.sources:
                        self.assertIn(source, notes.text)

    def test_real_screenshot_hashes_and_no_write_capture_policy(self) -> None:
        register = json.loads((DIRECTORY / "screenshots" / "captures.json").read_text(encoding="utf-8"))
        self.assertEqual(register["new_tests_submitted"], 0)
        self.assertFalse(register["non_read_api_requests_permitted"])
        self.assertEqual(len(register["images"]), 8)
        for item in register["images"]:
            with self.subTest(image=item["file"]):
                self.assertEqual(sha256(DIRECTORY / "screenshots" / item["file"]), item["sha256"])

    @unittest.skipUnless(Path(r"C:\Windows\Fonts\arial.ttf").is_file(), "Windows layout-metric fonts are required.")
    def test_text_fitting_honours_minimum_size_and_rejects_overflow(self) -> None:
        size, lines = fitted_lines("Preserve the raw text.", 4.0, 1.0, 22, 18, BODY_FONT)
        self.assertGreaterEqual(size, 18)
        self.assertTrue(lines)
        with self.assertRaises(ValueError):
            fitted_lines("This cannot fit.", 0.1, 0.1, 22, 18, BODY_FONT)

    def test_controlled_evidence_remains_the_documented_corpus(self) -> None:
        evidence = json.loads((ROOT / "docs" / "evidence" / "final-analysis-20260927.json").read_text(encoding="utf-8"))
        self.assertEqual(len(evidence["results"]), 12)
        self.assertEqual(evidence["statistics"]["total_tokens"], 81)
        self.assertEqual(sum(row["parse"]["accepted"] for row in evidence["results"]), 10)
        self.assertEqual(len(evidence["controls"]), 7)
        self.assertEqual(len(evidence["lexer_controls"]), 8)

    @unittest.skipUnless(Path(r"C:\Windows\Fonts\arial.ttf").is_file(), "Windows layout-metric fonts are required.")
    def test_code_indentation_is_preserved(self) -> None:
        self.assertEqual(wrap('{\n  "capabilities": {\n    "analyze": true\n  }\n}', 7, 18, BODY_FONT, False), [
            "{", '  "capabilities": {', '    "analyze": true', "  }", "}",
        ])

    def test_shapes_do_not_inherit_unintended_theme_effects(self) -> None:
        for output in (OUTPUT, BRIEF.output):
            with self.subTest(presentation=output.name):
                presentation = Presentation(str(output))
                for slide in presentation.slides:
                    for shape in slide.shapes:
                        for effect in shape.element.xpath(".//a:effectRef"):
                            self.assertEqual(effect.get("idx"), "0")


class BriefPresentationTests(unittest.TestCase):
    def test_twelve_essential_slides_with_source_references(self) -> None:
        specs = load_specs(BRIEF)
        self.assertEqual(len(specs), 12)
        self.assertEqual(specs[0].layout, "cover")
        self.assertEqual(specs[-1].layout, "closing")
        self.assertTrue(all(spec.sources and spec.notes for spec in specs))
        self.assertFalse(any("100-slide" in spec.notes for spec in specs))

    def test_brief_powerpoint_count_editability_and_notes(self) -> None:
        result = validate_pptx(BRIEF.output, 12)
        self.assertEqual(result["slides"], 12)
        self.assertEqual(result["speaker_note_pages"], 12)
        self.assertGreater(result["editable_text_shapes"], 100)
        self.assertGreaterEqual(result["picture_placements"], 6)
        self.assertEqual(result["editable_charts"], 1)
        self.assertEqual(result["off_canvas_shapes"], 0)
        presentation = Presentation(str(BRIEF.output))
        charts = [
            shape.chart for slide in presentation.slides for shape in slide.shapes
            if isinstance(shape, GraphicFrame) and shape.has_chart
        ]
        self.assertEqual(list(charts[0].series[0].values), [10.0, 2.0])

    def test_brief_content_matches_saved_powerpoint(self) -> None:
        presentation = Presentation(str(BRIEF.output))
        for number, (spec, slide) in enumerate(zip(load_specs(BRIEF), presentation.slides), 1):
            titles = [
                shape.text for shape in slide.shapes
                if isinstance(shape, Shape) and shape.name.startswith("title-")
            ]
            with self.subTest(slide=number):
                self.assertEqual(len(titles), 1)
                self.assertEqual(compact(titles[0]), compact(spec.title))
                self.assertNotIn(slide.element.get("show"), {"0", "false"})
                text = "\n".join(
                    shape.text for shape in slide.shapes
                    if isinstance(shape, Shape) and shape.has_text_frame
                )
                self.assertIn(f"{number:02d} / 12", text)
                self.assertNotIn(" / 100", text)
                self.assertNotIn("100-slide", text)
                source_links = {
                    shape.click_action.hyperlink.address for shape in slide.shapes
                    if shape.name.startswith("source-")
                }
                self.assertEqual(source_links, {"../" + source for source in spec.sources})
                notes = slide.notes_slide.notes_text_frame
                self.assertIsNotNone(notes)
                if notes is not None:
                    self.assertIn(spec.notes, notes.text)
                    self.assertIn(f"/ 12 - {spec.title}", notes.text)
                    for source in spec.sources:
                        self.assertIn(source, notes.text)
                        self.assertIn(compact(source), compact(text))

    def test_brief_cover_keeps_all_three_members(self) -> None:
        presentation = Presentation(str(BRIEF.output))
        text = "\n".join(
            shape.text for shape in presentation.slides[0].shapes
            if isinstance(shape, Shape) and shape.has_text_frame
        )
        for name, matricule in load_specs(BRIEF)[0].points:
            self.assertIn(name, text)
            self.assertIn(matricule, text)

    def test_brief_navigation_stays_within_the_short_deck(self) -> None:
        presentation = Presentation(str(BRIEF.output))
        navigation = load_specs(BRIEF)[2]
        actual = []
        for shape in presentation.slides[2].shapes:
            if shape.click_action.action == PP_ACTION.NAMED_SLIDE:
                destination = shape.click_action.target_slide
                if destination is None:
                    self.fail("A navigation card has no destination.")
                actual.append(destination.slide_id)
        expected = [
            presentation.slides[number - 1].slide_id
            for number in navigation.targets for _ in range(4)
        ]
        self.assertEqual(len(actual), 20)
        self.assertCountEqual(actual, expected)

    def test_brief_package_contains_the_original_visuals(self) -> None:
        expected = [
            DIRECTORY / "screenshots" / "analyzer-desktop.png",
            DIRECTORY / "visuals" / "analyzer-input.png",
            ROOT / "docs" / "evidence" / "screenshots" / "public-test-20260927.png",
            DIRECTORY / "visuals" / "dictionary-kass.png",
            DIRECTORY / "visuals" / "illustration-architecture.png",
            ROOT / "docs" / "diagrams" / "automaton-lexer.png",
            ROOT / "docs" / "diagrams" / "automaton-parser.png",
        ]
        with ZipFile(BRIEF.output) as archive:
            self.assertIsNone(archive.testzip())
            embedded = {
                hashlib.sha256(archive.read(name)).hexdigest()
                for name in archive.namelist() if name.startswith("ppt/media/")
            }
        self.assertEqual(embedded, {sha256(path) for path in expected})

    def test_brief_pdf_and_verification_match_the_published_files(self) -> None:
        manifest = json.loads(BRIEF.manifest.read_text(encoding="utf-8"))
        pdf_path = BRIEF.output.with_suffix(".pdf")
        self.assertEqual(manifest["slides"], 12)
        self.assertEqual(manifest["speaker_note_pages"], 12)
        self.assertEqual(manifest["pptx_sha256"], sha256(BRIEF.output))
        self.assertEqual(manifest["pdf_sha256"], sha256(pdf_path))
        self.assertTrue(manifest["rendered_pdf_verified"])
        audit = manifest["render_audit"]
        self.assertEqual(audit["pages"], 12)
        self.assertEqual(audit["text_boxes_checked"], manifest["editable_text_shapes"])
        self.assertEqual(audit["text_fit_failures"], [])
        self.assertEqual(audit["out_of_bounds_image_pages"], [])
        self.assertEqual(audit["unembedded_fonts"], [])
        with pymupdf.open(pdf_path) as document:
            self.assertEqual(len(document), 12)


if __name__ == "__main__":
    unittest.main()
