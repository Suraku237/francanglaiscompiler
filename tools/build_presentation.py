"""Build and verify the concise or full source-grounded Camfranglais presentation."""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from functools import lru_cache
import hashlib
import json
import math
from pathlib import Path
import re
import subprocess
from typing import Any
import unicodedata

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.slide import Slide
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / "docs" / "presentation"
OUTPUT = ROOT / "docs" / "camfranglais-presentation-100-slides.pptx"
REVIEW = ROOT / "docs" / ".build" / "presentation"
WIDTH, HEIGHT = 13.333333, 7.5
NAVY, TEAL, GOLD = "123041", "147E80", "CBA261"
INK, MUTED, PAPER = "183747", "587080", "F5F7F6"
MINT, LINE, WHITE = "E8F3EF", "D8E4E2", "FFFFFF"
RUST, PALE_GOLD = "A8583C", "FBF3E5"
BODY_FONT, DISPLAY_FONT, CODE_FONT = "Arial", "Georgia", "Consolas"
CHAPTERS = (
    "Meet the product", "Franc Analyzer", "Analysis and evidence", "Reference libraries",
    "Browser to API", "Inside the lexer", "Inside the grammar", "Automata and parsing",
    "Persistence and deployment", "Evidence and demonstration",
)
SHORT = {
    "FRENCH_FUNCTION_WORD": "FF", "ENGLISH_FUNCTION_WORD": "EF",
    "PIDGIN_MARKER": "PIDGIN", "Utterance_LF1": "Uf", "Utterance": "Ut",
    "SubjectTail": "St", "Clause": "Cl", "PredicateTail": "Pt",
    "ComplementTail": "Ct", "LinkedPhrase": "Lp",
}
FONT_FILES = {
    ("Arial", False): "arial.ttf", ("Arial", True): "arialbd.ttf",
    ("Georgia", False): "georgia.ttf", ("Georgia", True): "georgiab.ttf",
    ("Consolas", False): "consola.ttf", ("Consolas", True): "consolab.ttf",
}


@dataclass(frozen=True)
class Edition:
    count: int
    source: Path
    output: Path
    review: Path
    manifest: Path


REFERENCE = Edition(100, DIRECTORY / "slides.json", OUTPUT, REVIEW, DIRECTORY / "verification.json")
BRIEF = Edition(
    12, DIRECTORY / "slides-brief.json",
    ROOT / "docs" / "camfranglais-presentation-12-slides.pptx",
    REVIEW / "brief", DIRECTORY / "verification-brief.json",
)


def string(value: object, label: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{label} must be a string.")
    return value


def strings(value: object, label: str) -> list[str]:
    if not isinstance(value, list):
        raise ValueError(f"{label} must be a list.")
    return [string(item, label) for item in value]


@dataclass
class Spec:
    layout: str
    title: str
    subtitle: str
    sources: list[str]
    notes: str
    points: list[tuple[str, str]] = field(default_factory=list)
    image: str = ""
    second_image: str = ""
    art: str = ""
    code: str = ""
    data: str = ""
    before: str = ""
    after: str = ""
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    targets: list[int] = field(default_factory=list)


def parse_spec(value: object) -> Spec:
    if not isinstance(value, dict):
        raise ValueError("A slide must be an object.")
    points = value.get("points", [])
    rows = value.get("rows", [])
    targets = value.get("targets", [])
    if not isinstance(points, list) or not isinstance(rows, list):
        raise ValueError("Slide points and rows must be lists.")
    pairs = [strings(item, "point") for item in points]
    if any(len(pair) != 2 for pair in pairs):
        raise ValueError("Every point must have a heading and explanation.")
    if not isinstance(targets, list) or not all(isinstance(item, int) for item in targets):
        raise ValueError("Slide targets must be integer slide numbers.")
    return Spec(
        **{key: string(value.get(key, ""), key) for key in (
            "layout", "title", "subtitle", "notes", "image", "second_image", "art",
            "code", "data", "before", "after",
        )},
        sources=strings(value.get("sources"), "sources"),
        points=[(pair[0], pair[1]) for pair in pairs],
        columns=strings(value.get("columns", []), "columns"),
        rows=[strings(row, "row") for row in rows],
        targets=list(targets),
    )


def load_specs(edition: Edition = REFERENCE) -> list[Spec]:
    raw = json.loads(edition.source.read_text(encoding="utf-8"))
    if not isinstance(raw, list):
        raise ValueError("The presentation content must be a slide list.")
    specs = [parse_spec(item) for item in raw]
    if len(specs) != edition.count:
        raise ValueError(f"Exactly {edition.count} slides are required, not {len(specs)}.")
    if len({item.title for item in specs}) != edition.count:
        raise ValueError("All slide titles must be distinct.")
    for index, spec in enumerate(specs, 1):
        if not spec.notes or not spec.sources:
            raise ValueError(f"Slide {index} needs speaker notes and source references.")
        for source in spec.sources:
            if not (ROOT / source).is_file():
                raise ValueError(f"Slide {index} references a missing source: {source}")
        if any(not 1 <= target <= edition.count for target in spec.targets):
            raise ValueError(f"Slide {index} has an invalid navigation target.")
    return specs


@lru_cache(maxsize=128)
def font(name: str, size: float, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = Path(r"C:\Windows\Fonts") / FONT_FILES[(name, bold)]
    if not path.is_file():
        raise FileNotFoundError(f"Required layout-metric font is unavailable: {path}")
    return ImageFont.truetype(str(path), round(size * 4))


def wrap(text: str, width: float, size: float, name: str, bold: bool) -> list[str]:
    face = font(name, size, bold)
    limit = width * 72 * 4
    lines: list[str] = []
    for paragraph in text.split("\n"):
        indent = paragraph[:len(paragraph) - len(paragraph.lstrip(" "))]
        current = indent
        for word in paragraph[len(indent):].split(" "):
            candidate = f"{current} {word}" if current.strip() else current + word
            if face.getlength(candidate) <= limit:
                current = candidate
            else:
                if current.strip():
                    lines.append(current)
                if face.getlength(indent + word) > limit:
                    return []
                current = indent + word
        lines.append(current)
    return lines


def fitted_lines(
    text: str, width: float, height: float, size: float, minimum: float,
    name: str, bold: bool = False,
) -> tuple[float, list[str]]:
    candidate = size
    while candidate >= minimum:
        lines = wrap(text, width, candidate, name, bold)
        if lines and len(lines) * candidate * 1.16 <= height * 72 - 1:
            return candidate, lines
        candidate -= 0.5
    raise ValueError(f"Text does not fit above {minimum} pt in {width} x {height}: {text!r}")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def remove_theme_effects(shape: BaseShape) -> None:
    # Impress may still apply the theme's effectRef despite an empty effectLst.
    shape.shadow.inherit = False
    for reference in shape.element.xpath(".//a:effectRef"):
        reference.set("idx", "0")


def make_art(name: str, destination: Path) -> None:
    image = Image.new("RGB", (1500, 1050), "#" + NAVY)
    draw = ImageDraw.Draw(image)
    for x in range(40, 1500, 65):
        for y in range(35, 1050, 65):
            draw.ellipse((x, y, x + 3, y + 3), fill="#285060")
    teal, gold, light = "#43AAA7", "#" + GOLD, "#E4F2EE"
    title_font = ImageFont.truetype(str(Path(r"C:\Windows\Fonts\arialbd.ttf")), 38)
    small_font = ImageFont.truetype(str(Path(r"C:\Windows\Fonts\arial.ttf")), 29)
    if name in {"architecture", "storage"}:
        labels = ["REACT / UI", "FASTAPI / API", "PYTHON / RULES"] if name == "architecture" else [
            "PUBLIC POLICY", "SHARED WORKSPACE", "SQLITE / SNAPSHOTS",
        ]
        for index, label in reversed(list(enumerate(labels))):
            cy = 280 + index * 230
            draw.polygon([(240, cy), (690, cy - 140), (1210, cy), (760, cy + 140)], fill="#215061", outline=teal, width=4)
            draw.polygon([(240, cy), (760, cy + 140), (760, cy + 205), (240, cy + 65)], fill="#183C4B", outline=teal, width=3)
            draw.polygon([(760, cy + 140), (1210, cy), (1210, cy + 65), (760, cy + 205)], fill="#126565", outline=teal, width=3)
            draw.text((560, cy - 20), label, fill=light, font=title_font)
        draw.line([(760, 70), (760, 130)], fill=gold, width=6)
    else:
        centre = (755, 515)
        positions = [(420, 240), (1080, 240), (1180, 680), (745, 870), (330, 690)]
        if name == "tokens":
            labels = ["je / FF", "suis / VERB", "kass / ADJ", "go / VERB", "nang / VERB"]
            middle = "TOKENS"
        elif name == "grammar":
            labels, middle = ["RULES", "FIRST", "FOLLOW", "LL(1)", "STACK"], "GRAMMAR"
        else:
            labels, middle = ["FRENCH", "ENGLISH", "PIDGIN", "RAW TEXT", "CONTEXT"], "LANGUAGE"
        for i, (cx, cy) in enumerate(positions):
            draw.line([centre, (cx, cy)], fill="#398082", width=4)
            draw.rounded_rectangle((cx - 165, cy - 63, cx + 165, cy + 63), 24, fill="#1D4856", outline=gold if i == 3 else teal, width=4)
            length = draw.textlength(labels[i], font=small_font)
            draw.text((cx - length / 2, cy - 17), labels[i], fill=light, font=small_font)
        draw.ellipse((570, 330, 940, 700), fill="#147E80", outline=gold, width=6)
        length = draw.textlength(middle, font=title_font)
        draw.text((755 - length / 2, 494), middle, fill=light, font=title_font)
        for radius in [208, 220]:
            draw.arc((755 - radius, 515 - radius, 755 + radius, 515 + radius), 210, 330, fill=gold, width=3)
    image.save(destination, optimize=True)


def prepare_assets() -> dict[str, Path]:
    asset_dir = DIRECTORY / "visuals"
    asset_dir.mkdir(exist_ok=True)
    shots = DIRECTORY / "screenshots"
    assets = {name: shots / f"{name}.png" for name in (
        "analyzer-desktop", "analysis-desktop", "collection-desktop", "dictionary-desktop",
        "examples-desktop", "dictionary-search", "analyzer-mobile", "collection-mobile",
    )}
    originals = ROOT / "docs" / "evidence" / "screenshots"
    assets.update({
        "saved-token-panel": originals / "public-test-20260927.png",
        "lexer-dfa": ROOT / "docs" / "diagrams" / "automaton-lexer.png",
        "parser-pda": ROOT / "docs" / "diagrams" / "automaton-parser.png",
    })
    crops = [
        ("analyzer-input", "analyzer-desktop", (0.18, 0.10, 0.98, 0.68)),
        ("collection-filters", "collection-desktop", (0.18, 0.38, 0.98, 0.68)),
        ("dictionary-kass", "dictionary-search", (0.18, 0.50, 0.98, 0.95)),
    ]
    records: list[dict[str, object]] = []
    for target, original, fractions in crops:
        path = asset_dir / f"{target}.png"
        with Image.open(assets[original]) as image:
            w, h = image.size
            box = (
                round(fractions[0] * w), round(fractions[1] * h),
                round(fractions[2] * w), round(fractions[3] * h),
            )
            image.crop(box).save(path, optimize=True)
        assets[target] = path
        records.append({
            "file": path.relative_to(ROOT).as_posix(), "sha256": sha256(path),
            "derived_from": assets[original].relative_to(ROOT).as_posix(),
            "crop_box_pixels": box, "edit": "Crop only; no text, result or interface altered.",
        })
    for name in ("language", "tokens", "grammar", "architecture", "storage"):
        path = asset_dir / f"illustration-{name}.png"
        make_art(name, path)
        assets[f"art-{name}"] = path
        records.append({
            "file": path.relative_to(ROOT).as_posix(), "sha256": sha256(path),
            "origin": "Original programmatically drawn explanatory artwork; not a screenshot or fieldwork photograph.",
        })
    for key, path in assets.items():
        with Image.open(path) as image:
            image.verify()
        if not path.is_file():
            raise FileNotFoundError(f"Missing visual {key}: {path}")
    (asset_dir / "provenance.json").write_text(json.dumps(records, indent=2) + "\n", encoding="utf-8")
    return assets


class Canvas:
    def __init__(self, slide: Slide, number: int, dark: bool = False, total: int = 100) -> None:
        self.slide, self.number, self.dark = slide, number, dark
        self.total = total
        self.texts: list[dict[str, object]] = []
        self.layout_errors: list[str] = []
        self.background = NAVY if dark else PAPER
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(self.background)

    def box(
        self, x: float, y: float, w: float, h: float, fill: str,
        line: str | None = None, rounded: bool = False,
    ) -> Shape:
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h),
        )
        if rounded:
            shape.adjustments[0] = 0.11
        remove_theme_effects(shape)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line:
            shape.line.color.rgb = RGBColor.from_string(line)
            shape.line.width = Pt(0.7)
        else:
            shape.line.fill.background()
        return shape

    def circle(self, x: float, y: float, diameter: float, fill: str) -> Shape:
        shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
        remove_theme_effects(shape)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
        return shape

    def line(self, x1: float, y1: float, x2: float, y2: float, colour: str = LINE, width: float = 1) -> None:
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        remove_theme_effects(shape)
        shape.line.color.rgb = RGBColor.from_string(colour)
        shape.line.width = Pt(width)

    def text(
        self, value: str, x: float, y: float, w: float, h: float, size: float = 20,
        colour: str | None = None, bold: bool = False, name: str = BODY_FONT,
        minimum: float | None = None, align: PP_ALIGN = PP_ALIGN.LEFT,
        role: str = "body", middle: bool = False,
    ) -> Shape:
        minimum = size if minimum is None else minimum
        try:
            actual, lines = fitted_lines(value, w, h, size, minimum, name, bold)
        except ValueError as error:
            self.layout_errors.append(f"Slide {self.number:03d}: {error}")
            actual, lines = minimum, wrap(value, w, minimum, name, bold) or [value]
        shape = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(0)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE if middle else MSO_ANCHOR.TOP
        for index, value_line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = value_line
            paragraph.alignment = align
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = Pt(actual * 1.16)
            paragraph.font.name = name
            paragraph.font.size = Pt(actual)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(colour or (WHITE if self.dark else INK))
        shape.name = f"{role}-{len(self.texts) + 1}"
        self.texts.append({
            "shape": shape.name, "text": value, "role": role, "font": name, "font_points": actual,
            "box_inches": [x, y, w, h], "lines": len(lines),
            "predicted_height_inches": len(lines) * actual * 1.16 / 72,
        })
        return shape

    def picture(self, path: Path, x: float, y: float, w: float, h: float) -> None:
        with Image.open(path) as image:
            ratio = image.width / image.height
        out_w, out_h = (w, w / ratio) if w / h < ratio else (h * ratio, h)
        self.slide.shapes.add_picture(str(path), Inches(x + (w - out_w) / 2), Inches(y + (h - out_h) / 2), Inches(out_w), Inches(out_h))

    def browser(self, path: Path, x: float, y: float, w: float, h: float, label: str = "camfranglais.duckdns.org") -> None:
        self.box(x + 0.045, y + 0.06, w, h, "DCE5E4", rounded=True)
        self.box(x, y, w, h, WHITE, LINE, rounded=True)
        self.box(x, y, w, 0.32, MINT)
        for i, colour in enumerate((RUST, GOLD, TEAL)):
            self.circle(x + 0.14 + i * 0.15, y + 0.10, 0.08, colour)
        self.text(label, x + 0.74, y + 0.065, w - 1.1, 0.19, 10, MUTED, role="caption")
        self.picture(path, x + 0.06, y + 0.35, w - 0.12, h - 0.42)

    def header(self, spec: Spec) -> None:
        chapter = (self.number - 1) // 10
        section = "THE ESSENTIALS / CS4110" if self.total == 12 else f"{chapter + 1:02d} / {CHAPTERS[chapter].upper()}"
        self.text(section, 0.58, 0.29, 10.3, 0.23, 10.5, TEAL, True, role="section")
        self.text("CAMFRANGLAIS", 10.53, 0.29, 2.23, 0.23, 10.5, MUTED, True, align=PP_ALIGN.RIGHT, role="section")
        self.text(spec.title, 0.58, 0.80, 12.18, 0.63, 33, bold=True, minimum=28, role="title")
        self.text(spec.subtitle, 0.61, 1.54, 12.10, 0.46, 16.5, MUTED, minimum=16, role="subtitle")
        self.box(0.59, 1.36, 0.56, 0.035, GOLD)

    def footer(self, spec: Spec) -> None:
        colour = "B2CBCB" if self.dark else MUTED
        self.line(0.58, 6.63, 12.75, 6.63, "355565" if self.dark else LINE, 0.7)
        self.text("IMPLEMENTATION / EVIDENCE", 0.59, 6.69, 4.8, 0.18, 9.5, GOLD if self.dark else TEAL, True, role="footer")
        for i, source in enumerate(spec.sources):
            x, y = 0.59 + (i % 2) * 6.10, 6.93 + (i // 2) * 0.19
            text = self.text(source, x, y, 6.0, 0.19, 10.5, colour, name=CODE_FONT, minimum=10, role="source")
            text.click_action.hyperlink.address = "../" + source
        self.text("CS4110  /  SET A  /  27 SEP 2026", 0.59, 7.30, 8.0, 0.17, 9, colour, role="footer")
        self.text(f"{self.number:02d} / {self.total}", 11.35, 7.27, 1.40, 0.21, 11, colour, True, align=PP_ALIGN.RIGHT, role="footer")
        self.box(0, 7.46, WIDTH, 0.04, "284857" if self.dark else LINE)
        self.box(0, 7.46, WIDTH * self.number / self.total, 0.04, GOLD if self.dark else TEAL)


def points_column(canvas: Canvas, points: list[tuple[str, str]], x: float, y: float, w: float, h: float) -> None:
    step = h / len(points)
    for index, (heading, body) in enumerate(points):
        top = y + index * step
        canvas.box(x, top + 0.02, 0.045, min(0.54, step - 0.1), GOLD if canvas.dark else TEAL)
        canvas.text(heading, x + 0.20, top, w - 0.20, 0.42, 18.5, bold=True, minimum=17)
        canvas.text(body, x + 0.20, top + 0.48, w - 0.22, step - 0.55, 18, "BBD2D3" if canvas.dark else MUTED, minimum=17)


def cards(canvas: Canvas, spec: Spec, *, files: bool = False) -> None:
    count = len(spec.points)
    columns = 3 if count in (3, 5, 6) else 2
    rows = math.ceil(count / columns)
    gap, left, top, width, height = 0.22, 0.59, 2.22, 12.16, 4.03
    cw, ch = (width - gap * (columns - 1)) / columns, (height - gap * (rows - 1)) / rows
    for index, (heading, body) in enumerate(spec.points):
        x, y = left + (index % columns) * (cw + gap), top + (index // columns) * (ch + gap)
        canvas.box(x, y, cw, ch, WHITE, LINE, rounded=True)
        canvas.box(x, y, 0.055, ch, TEAL if index % 2 == 0 else GOLD)
        if files:
            path = Path(heading)
            canvas.text(path.parent.as_posix() if path.parent != Path(".") else "repository root", x + 0.23, y + 0.17, cw - 0.46, 0.22, 10.5, TEAL, name=CODE_FONT, role="caption")
            canvas.text(path.name, x + 0.23, y + 0.49, cw - 0.46, 0.42, 20, bold=True, minimum=17)
            canvas.text(body, x + 0.23, y + 1.01, cw - 0.46, ch - 1.14, 18, MUTED, minimum=16.5)
        elif rows == 1:
            canvas.circle(x + 0.25, y + 0.30, 0.66, MINT)
            canvas.text(f"{index + 1:02d}", x + 0.25, y + 0.46, 0.66, 0.32, 17, TEAL, True, align=PP_ALIGN.CENTER)
            canvas.text(heading, x + 0.28, y + 1.26, cw - 0.56, 0.75, 22, bold=True, minimum=19)
            canvas.text(body, x + 0.28, y + 2.13, cw - 0.56, ch - 2.30, 20, MUTED, minimum=18)
        else:
            canvas.text(heading, x + 0.26, y + 0.27, cw - 0.52, 0.45, 19.5, TEAL, True, minimum=18)
            canvas.text(body, x + 0.26, y + 0.87, cw - 0.52, ch - 1.02, 20, MUTED, minimum=18)


def flow(canvas: Canvas, spec: Spec) -> None:
    count, width, gap = len(spec.points), 12.15, 0.31
    cw = (width - (count - 1) * gap) / count
    for i, (heading, body) in enumerate(spec.points):
        x = 0.60 + i * (cw + gap)
        canvas.box(x, 2.70, cw, 2.81, WHITE, LINE, rounded=True)
        canvas.text(f"{i + 1:02d}", x + 0.20, 2.94, cw - 0.40, 0.53, 28, GOLD, True)
        canvas.text(heading, x + 0.20, 3.73, cw - 0.40, 0.60, 16.5, TEAL, True, minimum=15)
        canvas.text(body, x + 0.20, 4.42, cw - 0.40, 0.84, 18, MUTED, minimum=16.5)
        if i < count - 1:
            arrow = canvas.slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + cw + 0.085), Inches(3.85), Inches(0.14), Inches(0.24))
            remove_theme_effects(arrow)
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor.from_string(TEAL)
            arrow.line.fill.background()
    canvas.text("FOLLOW THE DATA  /  KEEP EACH RESPONSIBILITY VISIBLE", 0.61, 5.94, 12.0, 0.28, 11.5, MUTED, True, role="caption")


def draw_table(
    canvas: Canvas, columns: list[str], rows: list[list[str]],
    x: float = 0.60, y: float = 2.20, w: float = 12.13, h: float = 4.10,
    fractions: list[float] | None = None, font_size: float = 17.5,
) -> None:
    fractions = fractions or ([0.43, 0.57] if len(columns) == 2 else [0.25, 0.43, 0.32])
    if any(len(row) != len(columns) for row in rows):
        raise ValueError("Every table row must match its columns.")
    head_h = 0.47
    row_h = (h - head_h) / len(rows)
    canvas.box(x, y, w, head_h, TEAL)
    all_rows = [columns, *rows]
    for index, row in enumerate(all_rows):
        top = y if index == 0 else y + head_h + (index - 1) * row_h
        height = head_h if index == 0 else row_h
        if index:
            canvas.box(x, top, w, height, WHITE if index % 2 else MINT)
            canvas.line(x, top + height, x + w, top + height, LINE, 0.5)
        offset = x
        for j, value in enumerate(row):
            cell_w = w * fractions[j]
            canvas.text(value, offset + 0.16, top + 0.09, cell_w - 0.30, height - 0.13, 16 if index == 0 else font_size, WHITE if index == 0 else INK, index == 0, minimum=14.5 if index else 16)
            offset += cell_w


def abbreviate(value: str) -> str:
    for original in sorted(SHORT, key=len, reverse=True):
        value = value.replace(original, SHORT[original])
    return value


def content(canvas: Canvas, spec: Spec, assets: dict[str, Path], evidence: dict[str, Any]) -> list[tuple[Shape, int]]:
    links: list[tuple[Shape, int]] = []
    if spec.layout in {"cover", "closing"}:
        canvas.box(0.59, 0.63, 1.03, 0.055, GOLD)
        canvas.text("CAMFRANGLAIS / COMPILER WORKSPACE", 0.59, 0.99, 10.3, 0.29, 12, "85BABB", True, role="section")
        if spec.layout == "cover":
            canvas.text(spec.title, 0.59, 1.65, 6.6, 1.04, 47, WHITE, True, DISPLAY_FONT, minimum=42, role="title")
            canvas.text(spec.subtitle, 0.64, 2.95, 5.27, 1.52, 31, WHITE, name=BODY_FONT, minimum=29, role="subtitle")
            if canvas.total == 12:
                canvas.text("CS4110 / SUMMER 2026 / SET A", 0.65, 4.71, 5.50, 0.27, 12, GOLD, True)
                for i, (name, matricule) in enumerate(spec.points):
                    canvas.text(f"{name} / {matricule}", 0.65, 5.18 + i * 0.35, 5.66, 0.27, 12.5, "D1E2E1", minimum=12)
            else:
                canvas.text("THE APP. THE RULES. THE FILES.", 0.65, 4.83, 5.20, 0.33, 14, GOLD, True)
                canvas.text("100-slide technical presentation\nCS4110 / Summer 2026 / SET A", 0.65, 5.47, 5.0, 0.65, 16, "B2CBCB")
            canvas.browser(assets[spec.image], 6.63, 1.79, 6.08, 4.42)
        else:
            canvas.text(spec.title, 0.60, 1.86, 11.5, 1.32, 48, WHITE, name=DISPLAY_FONT, minimum=44, role="title")
            canvas.text(spec.subtitle, 0.65, 3.30, 11.5, 0.72, 31, "BBD2D3", role="subtitle")
            for i, (heading, body) in enumerate(spec.points):
                x = 0.66 + i * 4.10
                canvas.text(heading, x, 4.57, 3.50, 0.32, 13, GOLD, True)
                canvas.text(body, x, 5.02, 3.60, 0.45, 24, WHITE, minimum=22)
            site = canvas.text("camfranglais.duckdns.org", 0.66, 5.94, 10, 0.38, 19, "7DD2C7", True)
            site.click_action.hyperlink.address = "https://camfranglais.duckdns.org"
    elif spec.layout == "team":
        for i, (name, matricule) in enumerate(spec.points):
            x = 0.60 + i * 4.13
            canvas.box(x, 2.24, 3.89, 4.03, WHITE, LINE, True)
            canvas.circle(x + 0.29, 2.55, 0.90, MINT)
            canvas.text(f"0{i + 1}", x + 0.29, 2.81, 0.90, 0.42, 23, TEAL, True, align=PP_ALIGN.CENTER)
            canvas.text(name, x + 0.30, 3.83, 3.25, 1.33, 25, bold=True, minimum=23)
            canvas.text(matricule, x + 0.30, 5.55, 3.26, 0.40, 19, TEAL, True)
    elif spec.layout in {"cards", "files"}:
        cards(canvas, spec, files=spec.layout == "files")
    elif spec.layout == "navigation":
        for i, (heading, body) in enumerate(spec.points):
            cols, row = (i, 0) if i < 3 else (i - 3, 1)
            x, y = 0.60 + cols * 4.12, 2.21 + row * 2.08
            tile = canvas.box(x, y, 3.88, 1.86, WHITE, LINE, True)
            badge = canvas.text(f"0{i + 1}", x + 0.23, y + 0.19, 0.55, 0.33, 17, GOLD, True)
            label = canvas.text(heading, x + 0.23, y + 0.66, 3.40, 0.41, 22, TEAL, True, minimum=19)
            explanation = canvas.text(body, x + 0.23, y + 1.21, 3.40, 0.59, 17, MUTED, minimum=16.5)
            links.extend((shape, spec.targets[i]) for shape in (tile, badge, label, explanation))
        canvas.box(8.84, 4.29, 3.88, 1.86, NAVY, rounded=True)
        canvas.text("ONE SHARED\nPUBLIC WORKSPACE", 9.09, 4.68, 3.30, 1.10, 24, WHITE, True, minimum=22)
    elif spec.layout in {"flow", "deployment"}:
        flow(canvas, spec)
        if spec.layout == "deployment":
            canvas.text("DOCUMENTED TOPOLOGY / loopback upstream, public HTTPS", 0.61, 6.24, 12, 0.20, 10.5, TEAL, role="caption")
    elif spec.layout == "compare":
        for i, (heading, body) in enumerate(spec.points):
            x = 0.60 + i * 6.20
            canvas.box(x, 2.22, 5.94, 4.03, NAVY if i == 0 else WHITE, None if i == 0 else LINE, True)
            canvas.text("01" if i == 0 else "02", x + 0.32, 2.54, 1.3, 0.57, 31, GOLD, True)
            canvas.text(heading, x + 0.32, 3.40, 5.27, 0.71, 23, WHITE if i == 0 else TEAL, True, minimum=21)
            canvas.text(body, x + 0.32, 4.44, 5.24, 1.50, 22, "D1E2E1" if i == 0 else MUTED, minimum=19)
    elif spec.layout == "metrics":
        for i, (number, label) in enumerate(spec.points):
            x = 0.60 + (i % 2) * 6.19
            y = 2.24 + (i // 2) * 2.09
            canvas.box(x, y, 5.94, 1.84, WHITE, LINE, True)
            canvas.text(number, x + 0.27, y + 0.25, 5.34, 0.83, 48, TEAL, True, minimum=28)
            canvas.text(label, x + 0.28, y + 1.22, 5.33, 0.36, 19, MUTED, minimum=18)
    elif spec.layout in {"screenshot", "diagram"}:
        points_column(canvas, spec.points, 0.64, 2.29, 3.30, 3.91)
        if spec.layout == "screenshot":
            canvas.browser(assets[spec.image], 4.25, 2.19, 8.48, 4.08)
            canvas.text("REAL DEPLOYED APP / 27 SEP 2026 / crop only where noted", 4.29, 6.36, 8.25, 0.19, 9.5, MUTED, role="caption")
        else:
            canvas.box(4.23, 2.19, 8.50, 4.18, WHITE, LINE, True)
            canvas.picture(assets[spec.image], 4.38, 2.32, 8.20, 3.90)
    elif spec.layout == "diagram-wide":
        canvas.box(0.60, 2.15, 12.12, 3.24, WHITE, LINE, True)
        canvas.picture(assets[spec.image], 0.77, 2.29, 11.78, 2.96)
        for i, (heading, body) in enumerate(spec.points):
            x = 0.65 + i * 4.09
            canvas.text(heading.upper(), x, 5.59, 3.83, 0.28, 13, TEAL, True)
            canvas.text(body, x, 6.01, 3.86, 0.35, 18, MUTED, minimum=17)
    elif spec.layout == "art":
        canvas.box(5.55, 2.17, 7.16, 4.16, NAVY, rounded=True)
        canvas.picture(assets[f"art-{spec.art}"], 5.62, 2.24, 7.02, 4.02)
        points_column(canvas, spec.points, 0.65, 2.27, 4.37, 3.95)
    elif spec.layout == "mobile":
        points_column(canvas, spec.points, 0.65, 2.29, 4.04, 3.90)
        for i, asset in enumerate((spec.image, spec.second_image)):
            x = 5.47 + i * 3.52
            canvas.box(x, 2.15, 2.28, 4.28, NAVY, rounded=True)
            canvas.picture(assets[asset], x + 0.07, 2.27, 2.14, 3.99)
    elif spec.layout == "table":
        draw_table(canvas, spec.columns, spec.rows)
    elif spec.layout == "code":
        canvas.box(0.59, 2.18, 8.00, 4.15, NAVY, rounded=True)
        canvas.text("SOURCE WALKTHROUGH / ABBREVIATED WHERE MARKED", 0.86, 2.39, 7.46, 0.24, 10.5, GOLD, True, role="caption")
        canvas.text(spec.code, 0.87, 2.94, 7.42, 3.04, 19, "E0EFEB", name=CODE_FONT, minimum=16.5)
        points_column(canvas, spec.points, 8.97, 2.29, 3.51, 3.92)
    elif spec.layout == "rewrite":
        for top, label, value, colour in [(2.18, "BEFORE / ORIGINAL RULE", spec.before, WHITE), (3.78, "AFTER / TRANSFORMED RULES", spec.after, MINT)]:
            canvas.box(0.60, top, 12.12, 1.40, colour, LINE, True)
            canvas.text(label, 0.86, top + 0.19, 11.55, 0.24, 11.5, TEAL, True)
            canvas.text(value, 0.86, top + 0.59, 11.56, 0.66, 21, INK, name=CODE_FONT, minimum=18)
        for i, (heading, body) in enumerate(spec.points):
            x = 0.65 + i * 6.14
            canvas.text(heading.upper(), x, 5.55, 5.76, 0.29, 12, TEAL, True)
            canvas.text(body, x, 5.94, 5.77, 0.58, 18, MUTED, minimum=17)
    elif spec.layout == "set":
        canvas.box(0.60, 2.20, 12.13, 1.08, NAVY, rounded=True)
        canvas.text(spec.code, 0.85, 2.52, 11.60, 0.50, 30, WHITE, name=CODE_FONT, minimum=26)
        for i, (heading, body) in enumerate(spec.points):
            x = 0.60 + i * 4.12
            canvas.box(x, 3.67, 3.88, 2.54, WHITE, LINE, True)
            canvas.text(heading, x + 0.24, 3.99, 3.38, 0.64, 20, TEAL, True, minimum=15.5)
            canvas.text(body, x + 0.24, 4.98, 3.38, 1.03, 21, MUTED, minimum=19)
    elif spec.layout in {"priority", "route"}:
        step = 4.14 / len(spec.points)
        for i, (heading, body) in enumerate(spec.points):
            y = 2.18 + i * step
            canvas.box(0.60, y, 12.12, step - 0.06, WHITE if i % 2 == 0 else MINT, rounded=True)
            canvas.text(heading, 0.83, y + 0.17, 4.15, step - 0.23, 18, TEAL, True, minimum=16)
            canvas.text(body, 5.14, y + 0.17, 7.27, step - 0.22, 18, MUTED, minimum=16.5)
    elif spec.layout == "tokens":
        tokens = evidence["results"][1]["lexical"]["tokens"]
        canvas.text("je suis kass je go nang", 0.63, 2.23, 12, 0.70, 36, TEAL, name=DISPLAY_FONT, minimum=32)
        for i, token in enumerate(tokens):
            x = 0.62 + i * 2.035
            canvas.box(x, 3.27, 1.81, 1.35, WHITE, LINE, True)
            canvas.text(token["text"], x + 0.11, 3.48, 1.59, 0.44, 25, INK, True, align=PP_ALIGN.CENTER)
            canvas.text(abbreviate(token["category"]), x + 0.06, 4.05, 1.69, 0.28, 11.5, TEAL, True, align=PP_ALIGN.CENTER, minimum=10.5)
        for i, (heading, body) in enumerate(spec.points):
            x = 0.65 + i * 4.10
            canvas.text(heading, x, 5.16, 3.80, 0.39, 23, TEAL, True)
            canvas.text(body, x, 5.73, 3.82, 0.51, 18, MUTED, minimum=17)
    elif spec.layout == "bars":
        stats = evidence["statistics"]
        pairs = sorted(stats["category_counts"].items(), key=lambda pair: -pair[1]) if spec.data == "categories" else [
            (item["token"], item["count"]) for item in stats["frequencies"][:6]
        ]
        maximum, step = max(value for _, value in pairs), 3.83 / len(pairs)
        for i, (label, value) in enumerate(pairs):
            y = 2.24 + i * step
            canvas.text(abbreviate(label), 0.64, y + 0.045, 2.95, step - 0.03, 16.5, MUTED, minimum=15.5)
            canvas.box(3.59, y + 0.05, 8.24, step * 0.65, "E4EBE8", rounded=True)
            canvas.box(3.59, y + 0.05, 8.24 * value / maximum, step * 0.65, TEAL if i % 3 else "B89A63", rounded=True)
            canvas.text(str(value), 12.03, y + 0.035, 0.65, step - 0.02, 18, INK, True, minimum=17)
        canvas.text("CONTROLLED STUDY / FF = French function word; EF = English function word", 0.65, 6.22, 12, 0.24, 11.5, MUTED, role="caption")
    elif spec.layout in {"trace", "stack"}:
        trace = evidence["results"][1]["parse"]["trace"]
        if spec.layout == "trace":
            rows = []
            for item in trace[:4]:
                rows.append([
                    "[" + ", ".join(abbreviate(s) for s in item["stack"]) + "]",
                    " ".join(abbreviate(s) for s in item["remaining"]),
                    abbreviate(item["action"]),
                ])
            draw_table(canvas, ["Stack (top at right)", "Remaining input", "Action"], rows, h=3.31, fractions=[0.28, 0.33, 0.39], font_size=17)
            canvas.text(spec.points[0][1], 0.66, 5.80, 12.0, 0.49, 20, TEAL, True, minimum=18)
            canvas.text("Ut = Utterance / Uf = ending helper / Cl = Clause / St = SubjectTail / FF = French function word", 0.66, 6.30, 12.0, 0.19, 10.5, MUTED, role="caption")
        else:
            indices = (0, 1) if spec.data == "stack-expand" else (2, 3)
            for j, index in enumerate(indices):
                row = trace[index]
                y = 2.49 + j * 1.75
                canvas.text("BEFORE" if j == 0 else "AFTER", 0.65, y + 0.28, 1.23, 0.34, 13, TEAL, True)
                for i, symbol in enumerate(row["stack"]):
                    x = 2.00 + i * 2.48
                    canvas.box(x, y, 2.24, 0.99, NAVY if i == len(row["stack"]) - 1 else WHITE, LINE, True)
                    canvas.text(abbreviate(symbol), x + 0.10, y + 0.27, 2.04, 0.46, 26, WHITE if i == len(row["stack"]) - 1 else INK, True, align=PP_ALIGN.CENTER)
                canvas.text("TOP", 10.27, y + 0.30, 2.0, 0.35, 13, GOLD, True)
            canvas.text(spec.points[2][1], 0.67, 6.00, 12.0, 0.39, 20, TEAL, True, minimum=18)
    elif spec.layout == "controls":
        rows = [
            [item["id"], item["text"] or "(empty)", "ACCEPT" if item["accepted"] else "REJECT"]
            for item in evidence["controls"]
        ]
        draw_table(canvas, ["Control", "Constructed input", "Saved-grammar result"], rows, fractions=[0.17, 0.50, 0.33], font_size=18)
    elif spec.layout == "outcome":
        accepted = sum(item["parse"]["accepted"] for item in evidence["results"])
        rejected = len(evidence["results"]) - accepted
        data = ChartData()
        data.categories = ["Accepted", "Rejected"]
        data.add_series("Original statements", [accepted, rejected])
        container: object = canvas.slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT, Inches(0.65), Inches(2.14), Inches(6.24), Inches(4.15), data,
        )
        if not isinstance(container, GraphicFrame):
            raise TypeError("The chart insertion did not return a PowerPoint graphic frame.")
        chart = container.chart
        chart.has_legend = False
        chart.has_title = False
        for point, colour in zip(chart.series[0].points, (TEAL, RUST)):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = RGBColor.from_string(colour)
            point.format.line.fill.background()
        canvas.text(f"{accepted / (accepted + rejected):.1%}", 2.28, 3.71, 2.94, 0.58, 36, TEAL, True, align=PP_ALIGN.CENTER, minimum=33)
        points_column(canvas, spec.points, 7.37, 2.35, 5.00, 3.84)
    else:
        raise ValueError(f"Unknown slide layout: {spec.layout}")
    return links


def validate_pptx(path: Path, expected_count: int = 100) -> dict[str, int]:
    presentation = Presentation(str(path))
    if len(presentation.slides) != expected_count:
        raise ValueError(f"The published presentation does not contain exactly {expected_count} slides.")
    slide_width, slide_height = presentation.slide_width, presentation.slide_height
    if slide_width is None or slide_height is None:
        raise ValueError("The presentation is missing its slide dimensions.")
    editable, pictures, charts, note_count = 0, 0, 0, 0
    for number, slide in enumerate(presentation.slides, 1):
        notes = slide.notes_slide.notes_text_frame
        if notes is None or not notes.text.strip():
            raise ValueError(f"Slide {number} is missing notes.")
        note_count += 1
        for shape in slide.shapes:
            if shape.left < -10 or shape.top < -10 or shape.left + shape.width > slide_width + 10 or shape.top + shape.height > slide_height + 10:
                raise ValueError(f"Slide {number} has an off-canvas shape: {shape.name}")
            if isinstance(shape, Shape) and shape.has_text_frame and shape.text.strip():
                editable += 1
            if shape.shape_type == 13:
                pictures += 1
            if shape.has_chart:
                charts += 1
    return {"slides": expected_count, "speaker_note_pages": note_count, "editable_text_shapes": editable, "picture_placements": pictures, "editable_charts": charts, "off_canvas_shapes": 0}


def build(edition: Edition = REFERENCE) -> dict[str, object]:
    specs = load_specs(edition)
    assets = prepare_assets()
    evidence = json.loads((ROOT / "docs" / "evidence" / "final-analysis-20260927.json").read_text(encoding="utf-8"))
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(WIDTH), Inches(HEIGHT)
    properties = presentation.core_properties
    properties.title = f"Camfranglais | Application and Source Guide | {edition.count} Slides"
    properties.subject = "CS4110 SET A: public app, responsible source files, lexer and LL(1) parser"
    properties.author = "Kwete Ngouba Junior Rayan; Djemtchimo Noukui Bruno Jonatan; Amina Boubakary"
    properties.keywords = f"Camfranglais, compiler, CS4110, LL(1), DFA, source guide, {edition.count} slides"
    properties.comments = "Editable slides with actual app screenshots, source references and speaker notes."
    layouts = []
    layout_errors: list[str] = []
    links: list[tuple[Shape, int]] = []
    for number, spec in enumerate(specs, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        canvas = Canvas(slide, number, dark=spec.layout in {"cover", "closing"}, total=edition.count)
        if spec.layout not in {"cover", "closing"}:
            canvas.header(spec)
        links.extend(content(canvas, spec, assets, evidence))
        canvas.footer(spec)
        layout_errors.extend(canvas.layout_errors)
        notes = slide.notes_slide.notes_text_frame
        if notes is None:
            raise RuntimeError("The presentation template has no speaker-notes placeholder.")
        notes.text = (
            f"SLIDE {number:02d} / {edition.count} - {spec.title}\n\n{spec.notes}\n\n"
            "RESPONSIBLE FILES / EVIDENCE\n" + "\n".join(spec.sources)
            + "\n\nVisuals: real app screenshots or clearly labelled original explanatory artwork. "
            "Source paths are repository-relative. No application or VPS change is performed by this presentation."
        )
        layouts.append({"number": number, "title": spec.title, "layout": spec.layout, "text_boxes": canvas.texts})
    if layout_errors:
        raise ValueError("Presentation was not published because of layout failures:\n" + "\n".join(layout_errors))
    for shape, target in links:
        shape.click_action.target_slide = presentation.slides[target - 1]
    edition.review.mkdir(parents=True, exist_ok=True)
    staged = edition.review / edition.output.name
    presentation.save(str(staged))
    metrics = validate_pptx(staged, edition.count)
    edition.output.write_bytes(staged.read_bytes())
    revision = subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT, check=True, capture_output=True, text=True).stdout.strip()
    manifest = {
        "format": 1, "title": properties.title, "source_revision_at_build": revision,
        "application_baseline": evidence["source_revision"], "pptx": edition.output.relative_to(ROOT).as_posix(),
        "pptx_sha256": sha256(edition.output), "pptx_bytes": edition.output.stat().st_size,
        **metrics, "unique_layouts": len({s.layout for s in specs}),
        "referenced_source_files": sorted({source for spec in specs for source in spec.sources}),
        "authentic_screenshot_register": "docs/presentation/screenshots/captures.json",
        "derived_visual_register": "docs/presentation/visuals/provenance.json",
        "content_source": edition.source.relative_to(ROOT).as_posix(),
        "generation": "python -m tools.build_presentation" + (" --brief" if edition == BRIEF else ""),
        "rendered_pdf_verified": False,
    }
    (edition.review / "layout-audit.json").write_text(json.dumps(layouts, indent=2) + "\n", encoding="utf-8")
    edition.manifest.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return manifest


def compact(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text)).replace("\u00ad", "")


def verify_render(edition: Edition = REFERENCE) -> dict[str, object]:
    import pymupdf

    pdf_path = edition.review / edition.output.with_suffix(".pdf").name
    audit = json.loads((edition.review / "layout-audit.json").read_text(encoding="utf-8"))
    manifest_path = edition.manifest
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest["pptx_sha256"] != sha256(edition.output):
        raise ValueError("The PPTX changed after its layout audit; rebuild before verifying.")
    if pdf_path.stat().st_mtime < edition.output.stat().st_mtime:
        raise ValueError("The PDF predates the current PPTX. Render the current presentation first.")
    issues = []
    image_bounds = []
    with pymupdf.open(pdf_path) as pdf:
        if len(pdf) != edition.count:
            raise ValueError(f"The rendered deck contains {len(pdf)} pages instead of {edition.count}.")
        if len(audit) != edition.count:
            raise ValueError("The layout audit does not match the selected presentation edition.")
        pages = [pdf[index] for index in range(len(pdf))]
        fonts = {entry[0] for page in pages for entry in page.get_fonts(full=True)}
        unembedded = []
        for xref in sorted(fonts):
            font_record: object = pdf.extract_font(xref, named=True)
            if not isinstance(font_record, dict):
                raise TypeError("The PDF font extractor did not return the requested named record.")
            if not font_record.get("content"):
                unembedded.append(xref)
        for layout, page in zip(audit, pages):
            for item in layout["text_boxes"]:
                x, y, w, h = item["box_inches"]
                rectangle = pymupdf.Rect(x * 72 - 2, y * 72 - 2, (x + w) * 72 + 2, (y + h) * 72 + 2)
                extracted = page.get_text("text", clip=rectangle, sort=True)
                if not isinstance(extracted, str):
                    raise TypeError("Text extraction did not return plain text.")
                actual = compact(extracted)
                expected = compact(item["text"])
                if expected not in actual:
                    issues.append({
                        "slide": layout["number"], "shape": item["shape"],
                        "expected": item["text"], "observed": actual,
                    })
            for image in page.get_image_info():
                bounds = pymupdf.Rect(image["bbox"])
                if not (page.rect + (-0.2, -0.2, 0.2, 0.2)).contains(bounds):
                    image_bounds.append(layout["number"])
        report = {
            "pages": len(pdf), "text_boxes_checked": sum(len(page["text_boxes"]) for page in audit),
            "text_fit_failures": issues, "out_of_bounds_image_pages": image_bounds,
            "font_resources": len(fonts), "unembedded_fonts": unembedded,
        }
        (edition.review / "render-audit.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
        if issues or image_bounds or unembedded:
            raise ValueError(
                f"Rendered-slide validation failed: {len(issues)} text boxes, "
                f"{len(image_bounds)} image bounds, {len(unembedded)} unembedded fonts. "
                f"See {edition.review / 'render-audit.json'}."
            )
        review_dir = edition.review / "review"
        review_dir.mkdir(exist_ok=True)
        label_font = ImageFont.truetype(str(Path(r"C:\Windows\Fonts\arialbd.ttf")), 18)
        for start in range(0, edition.count, 16):
            sheet = Image.new("RGB", (1600, 1040), "#D8E4E2")
            draw = ImageDraw.Draw(sheet)
            for offset, number in enumerate(range(start, min(start + 16, edition.count))):
                pixmap = pdf[number].get_pixmap(matrix=pymupdf.Matrix(0.405, 0.405), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                x, y = (offset % 4) * 400, (offset // 4) * 260
                sheet.paste(image, (x + 5, y + 30))
                draw.text((x + 8, y + 5), f"{number + 1:03d}", font=label_font, fill="#" + INK)
            sheet.save(review_dir / f"contact-{start // 16 + 1:02d}.png")
        review_pages = list(range(1, 13)) if edition == BRIEF else [
            1, 2, 5, 11, 12, 17, 25, 36, 47, 49, 56, 65, 71, 73, 74, 77, 88, 92, 99, 100,
        ]
        for number in review_pages:
            pdf[number - 1].get_pixmap(matrix=pymupdf.Matrix(1.4, 1.4), alpha=False).save(review_dir / f"slide-{number:03d}.png")
    destination = edition.output.with_suffix(".pdf")
    destination.write_bytes(pdf_path.read_bytes())
    manifest.update({
        "rendered_pdf_verified": True, "pdf": destination.relative_to(ROOT).as_posix(),
        "pdf_sha256": sha256(destination), "pdf_bytes": destination.stat().st_size,
        "renderer": "Locally extracted LibreOffice 26.2.6.3; no cloud upload",
        "renderer_msi_sha256": "f9877032fd908beb9c0ddf06df4af5c2e85f419c42e14876c4cce5aae5fb2660",
        "render_audit": report,
    })
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Check the existing PPTX without regenerating it.")
    parser.add_argument("--verify-render", action="store_true", help="Validate the current locally rendered PDF and publish it.")
    parser.add_argument("--render", action="store_true", help="Build, render with local LibreOffice, and verify both formats.")
    parser.add_argument("--brief", action="store_true", help="Use the concise 12-slide presentation instead of the full reference deck.")
    args = parser.parse_args()
    edition = BRIEF if args.brief else REFERENCE
    if args.verify_render:
        print(json.dumps(verify_render(edition), indent=2))
    elif args.check:
        load_specs(edition)
        print(json.dumps(validate_pptx(edition.output, edition.count), indent=2))
    else:
        manifest = build(edition)
        if args.render:
            renderer = ROOT / "docs" / ".tools" / "LibreOffice-26.2.6" / "program" / "soffice.com"
            if not renderer.is_file():
                raise FileNotFoundError("The optional local LibreOffice renderer is missing; see docs/README.md.")
            profile = (edition.review / "libreoffice-profile").as_uri()
            subprocess.run([
                str(renderer), f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                "--outdir", str(edition.review), str(edition.output),
            ], check=True, timeout=900)
            verify_render(edition)
            manifest = json.loads(edition.manifest.read_text(encoding="utf-8"))
        print(json.dumps({key: value for key, value in manifest.items() if key != "referenced_source_files"}, indent=2))


if __name__ == "__main__":
    main()
