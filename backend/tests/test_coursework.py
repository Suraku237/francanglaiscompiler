import base64
import io
import json
import subprocess
import sys
import unittest
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation

from backend import coursework_store
from backend.tests.test_api import ApiTestCase
from compiler.parser.service import DEFAULT_GRAMMAR
from data_collector import dataset


class CourseworkTests(ApiTestCase):
    def setUp(self):
        super().setUp()
        self.profile = {
            "group_members": ["", "", ""], "grammar": DEFAULT_GRAMMAR,
            "manual_transcription_confirmed": False, "grammar_rationale": "", "discussion": "",
            "collection_method": "", "limitations": "",
        }

    def add_entry(self, text="Le taxi don refuse.", **fields):
        response = self.client.post("/api/dataset", json={"text": text, **fields})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def test_empty_workspace_is_honest_and_does_not_invent_data(self):
        response = self.client.get("/api/coursework")
        self.assertEqual(response.status_code, 200, response.text)
        state = response.json()
        self.assertEqual(state["stats"]["total"], 0)
        self.assertEqual(state["stats"]["sentences"], 0)
        self.assertEqual(state["project"]["group_members"], ["", "", ""])
        self.assertEqual(state["brief"]["report_pages"], [25, 30])
        self.assertEqual(len(state["stats"]["missing_topics"]), 10)
        statuses = {row["id"]: row["status"] for row in state["requirements"]}
        self.assertEqual(statuses["data"], "needs_input")
        self.assertEqual(statuses["report"], "review")
        self.assert_no_outbound_http()
        self.assertFalse((self.directory / "coursework" / "project.json").exists())

    def test_profile_and_grammar_persist_and_bad_updates_do_not_replace(self):
        self.profile.update(group_members=["Alice", "Bob", "Chantal"], grammar="S -> NOUN")
        response = self.client.put("/api/coursework/project", json=self.profile)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(coursework_store.load_project().grammar, "S -> NOUN")
        invalid = self.client.put("/api/coursework/project", json={**self.profile, "grammar": "S -> Misspelled"})
        self.assertEqual(invalid.status_code, 422)
        self.assertEqual(coursework_store.load_project().grammar, "S -> NOUN")

    def test_group_bounds_and_duplicate_names(self):
        for members in (["A"], ["A", "B", "C", "D"], ["A", "a", "C"]):
            with self.subTest(members=members):
                response = self.client.put("/api/coursework/project", json={**self.profile, "group_members": members})
                self.assertEqual(response.status_code, 422)

    def test_full_collection_evidence_counts_sentences_only(self):
        for index, category in enumerate(dataset.CATEGORIES[:-1]):
            self.add_entry(f"Statement {index}", category=category)
        self.add_entry("Extra word", entry_type="Word")
        self.profile.update(
            group_members=["Alice", "Bob", "Chantal"],
            manual_transcription_confirmed=True, collection_method="Manual field listening, examples described by group.",
        )
        self.client.put("/api/coursework/project", json=self.profile)
        state = self.client.get("/api/coursework").json()
        self.assertEqual(state["stats"]["sentences"], 10)
        self.assertEqual(state["stats"]["total"], 11)
        self.assertEqual(state["stats"]["missing_topics"], [])
        self.assertEqual(next(row for row in state["requirements"] if row["id"] == "data")["status"], "ready")

    def test_analyzer_uses_own_entries_and_calculates_variation(self):
        first = self.add_entry("Le taxi don refuse.")
        second = self.add_entry("le TAXI don refuse !")
        result = self.client.post("/api/coursework/analyze", json={"grammar": DEFAULT_GRAMMAR})
        self.assertEqual(result.status_code, 200, result.text)
        body = result.json()
        self.assertEqual(body["summary"]["total"], 2)
        self.assertEqual({test["id"] for test in body["tests"]}, {first["id"], second["id"]})
        self.assertEqual(body["lexical"]["total_tokens"], 10)
        self.assertIn({"token": "taxi", "count": 2}, body["lexical"]["frequencies"])
        self.assertTrue(any(row["normalized"] == "taxi" for row in body["lexical"]["variations"]))
        self.assert_no_outbound_http()

    def test_slang_annotations_respect_word_boundaries_and_preserve_source_spelling(self):
        cases = (
            ("pre-je wanda", [], ["pre-je", "wanda"]),
            ("je wanda-post", [], ["je", "wanda-post"]),
            ("JE\tWANDA", ["JE\tWANDA"], ["JE", "WANDA"]),
            ("Je   Wanda, JE\tWANDA!", ["Je   Wanda", "JE\tWANDA"], ["Je", "Wanda", ",", "JE", "WANDA", "!"]),
        )
        expected = [(self.add_entry(text)["id"], text, phrases, tokens) for text, phrases, tokens in cases]
        response = self.client.post("/api/coursework/analyze", json={"grammar": DEFAULT_GRAMMAR})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["summary"]["total"], len(cases))
        statements = {row["id"]: row for row in response.json()["lexical"]["statements"]}
        for entry_id, text, phrases, tokens in expected:
            with self.subTest(text=text):
                self.assertEqual(statements[entry_id]["text"], text)
                self.assertEqual(statements[entry_id]["slang_expressions"], phrases)
                self.assertEqual([token["text"] for token in statements[entry_id]["tokens"]], tokens)
        self.assert_no_outbound_http()

    def test_parse_trace_full_input_and_unsupported_symbol(self):
        good = self.client.post("/api/coursework/parse", json={"grammar": "S -> NOUN", "text": "taxi"}).json()
        bad = self.client.post("/api/coursework/parse", json={"grammar": "S -> NOUN", "text": "taxi +"}).json()
        self.assertTrue(good["parse"]["accepted"])
        self.assertTrue(good["parse"]["trace"])
        self.assertFalse(bad["parse"]["accepted"])
        self.assertEqual(bad["tokens"][-1], {"text": "+", "category": "UNKNOWN"})
        epsilon = self.client.post("/api/coursework/parse", json={"grammar": "S -> epsilon", "text": ""}).json()
        self.assertTrue(epsilon["parse"]["accepted"])

    def test_conflicts_are_reported_not_resolved(self):
        grammar = "S -> A | B\nA -> NOUN\nB -> NOUN"
        response = self.client.post("/api/coursework/analyze", json={"grammar": grammar})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertFalse(response.json()["grammar"]["is_ll1"])
        parsed = self.client.post("/api/coursework/parse", json={"grammar": grammar, "text": "taxi"}).json()
        self.assertFalse(parsed["parse"]["accepted"])
        self.assertTrue(parsed["parse"]["error"])

    def test_oversize_and_malformed_grammar_validation(self):
        for grammar in ("", "S ->", "S -> MISSING", "x" * 12001):
            self.assertEqual(self.client.post("/api/coursework/analyze", json={"grammar": grammar}).status_code, 422)

    def test_ai_explanation_is_removed_and_never_sends_collection(self):
        self.add_entry("PRIVATE FIELD OBSERVATION", contributor="PRIVATE CONTRIBUTOR")
        response = self.client.post("/api/coursework/explain", json={
            "grammar": "S -> NOUN", "text": "taxi", "question": "Explain why it is accepted.", "language": "en",
        })
        self.assertEqual(response.status_code, 404, response.text)
        self.assertNotIn("/api/coursework/explain", self.client.get("/openapi.json").json()["paths"])
        self.assert_no_outbound_http()

    @staticmethod
    def image_data():
        content = io.BytesIO()
        Image.new("RGB", (100, 60), color="green").save(content, format="PNG")
        return "data:image/png;base64," + base64.b64encode(content.getvalue()).decode()

    def test_screenshot_upload_read_delete_and_validation(self):
        response = self.client.post("/api/coursework/screenshots", json={"name": "Working parser", "data_url": self.image_data()})
        self.assertEqual(response.status_code, 201, response.text)
        saved = response.json()
        fetched = self.client.get(saved["url"])
        self.assertEqual(fetched.status_code, 200)
        self.assertEqual(fetched.headers["content-type"], "image/png")
        screenshots = self.client.get("/api/coursework").json()["screenshots"]
        self.assertEqual(len(screenshots), 1)
        self.assertEqual(screenshots[0]["name"], "Working parser")
        self.assertEqual(self.client.delete(saved["url"]).status_code, 204)
        self.assertEqual(self.client.get(saved["url"]).status_code, 404)
        self.assertEqual(self.client.get("/api/coursework/screenshots/not-a-valid-id").status_code, 404)
        for value in ("data:image/png;base64,bm90YW5pbWFnZQ==", "data:text/html;base64,SGk="):
            self.assertEqual(self.client.post("/api/coursework/screenshots", json={"name": "invalid", "data_url": value}).status_code, 422)

    def test_export_contains_real_results_report_slides_source_not_secrets(self):
        self.add_entry("taxi")
        self.profile.update(grammar="S -> NOUN", discussion="<script>do not execute</script>")
        self.client.put("/api/coursework/project", json=self.profile)
        self.client.post("/api/coursework/screenshots", json={"name": "Example image", "data_url": self.image_data()})
        response = self.client.get("/api/coursework/export")
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else "")
        self.assertEqual(response.headers["content-type"], "application/zip")
        with ZipFile(io.BytesIO(response.content)) as archive:
            names = archive.namelist()
            self.assertFalse(any(".env" in name or "node_modules" in name for name in names))
            self.assertIn("source/compiler/parser/service.py", names)
            self.assertIn("source/tests/test_collected.py", names)
            self.assertIn("source/dictionary/camfranglais.md", names)
            self.assertIn("source/dictionary/extra_lexicon.md", names)
            self.assertIn("source/examples/camfranglais_statements.csv", names)
            self.assertIn(b"Constructed example", archive.read("source/examples/camfranglais_statements.csv"))
            self.assertIn(b"**tchop** | to eat", archive.read("source/dictionary/camfranglais.md"))
            self.assertIn(b"**motard** | a motorcycle taxi rider", archive.read("source/dictionary/extra_lexicon.md"))
            report = archive.read("report.html").decode()
            self.assertEqual(report.count("<section class='report-page'>"), 25)
            self.assertIn("&lt;script&gt;", report)
            self.assertNotIn("<script>", report)
            self.assertIn("data:image/png;base64,", report)
            slides = Presentation(io.BytesIO(archive.read("presentation.pptx")))
            self.assertEqual(len(slides.slides), 10)
            notes = slides.slides[-1].notes_slide.notes_text_frame
            assert notes is not None
            self.assertIn("Minute 10 of 10", notes.text)
            cases = json.loads(archive.read("source/tests/collected_cases.json"))
            self.assertEqual(cases[0]["text"], "taxi")
            extracted = self.directory / "exported"
            archive.extractall(extracted)
        completed = subprocess.run(
            [sys.executable, "-m", "unittest", "discover", "-s", "tests", "-v"],
            cwd=extracted / "source", capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)

    def test_exporting_empty_data_is_explicit_draft(self):
        self.assertEqual(self.client.get("/api/dictionary").json()["total"], 179)
        self.assertEqual(self.client.get("/api/examples").json()["total"], 26)
        self.assertEqual(self.client.get("/api/coursework").json()["stats"]["total"], 0)
        response = self.client.get("/api/coursework/export")
        self.assertEqual(response.status_code, 200)
        with ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(json.loads(archive.read("source/tests/collected_cases.json")), [])
            self.assertIn("source/examples/camfranglais_statements.csv", archive.namelist())
            self.assertIn("TO COMPLETE", archive.read("report.html").decode())

    def test_learned_lexicon_matches_manual_corpus_and_exported_regression_results(self):
        entry = self.add_entry("zandolo", entry_type="Word", lexical_category="NOUN", review_status="approved")
        sentence = self.add_entry("zandolo zandolo", entry_type="Sentence")
        self.profile.update(grammar="S -> NOUN | NOUN NOUN")
        self.client.put("/api/coursework/project", json=self.profile)
        manual = self.client.post("/api/coursework/parse", json={
            "grammar": self.profile["grammar"], "text": "zandolo",
        }).json()
        self.assertEqual(manual["tokens"], [{"text": "zandolo", "category": "NOUN"}])
        self.assertTrue(manual["parse"]["accepted"])
        corpus = self.client.post("/api/coursework/analyze", json={"grammar": self.profile["grammar"]}).json()
        self.assertEqual(corpus["summary"]["accepted"], 2)
        self.assertEqual({case["id"] for case in corpus["tests"]}, {entry["id"], sentence["id"]})
        exported = self.client.get("/api/coursework/export")
        self.assertEqual(exported.status_code, 200)
        with ZipFile(io.BytesIO(exported.content)) as archive:
            root = self.directory / "learned-export"
            archive.extractall(root)
            token_csv = archive.read("artifacts/tokens.csv").decode()
            self.assertNotIn("UNKNOWN", token_csv)
        completed = subprocess.run(
            [sys.executable, "-m", "unittest", "discover", "-s", "tests", "-q"],
            cwd=root / "source", capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)


if __name__ == "__main__":
    unittest.main()
