import base64
import csv
import html
import io
import json
from collections.abc import Sequence
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.util import Inches, Pt

from data_collector import dataset

from .coursework import BRIEF, analyze_coursework, corpus_stats, lexical_spec, read_corpus
from .coursework_models import ProjectProfile
from .coursework_store import list_screenshots, load_project, screenshot_path
from .dictionary import DICTIONARY_PATHS
from .examples import EXAMPLES_PATH

ROOT = Path(__file__).resolve().parents[1]


def csv_bytes(headers: list[str], rows: list[dict]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2)


def _escape(value: object) -> str:
    return html.escape(str(value), quote=True)


def _paragraph(value: str, missing: str) -> str:
    return f"<p>{_escape(value or missing).replace(chr(10), '<br>')}</p>"


def _table(headers: list[str], rows: Sequence[Sequence[object]], maximum: int = 60) -> str:
    table = "<table><thead><tr>" + "".join(f"<th>{_escape(header)}</th>" for header in headers) + "</tr></thead><tbody>"
    table += "".join(
        "<tr>" + "".join(f"<td>{_escape(cell)}</td>" for cell in row) + "</tr>"
        for row in rows[:maximum]
    )
    table += "</tbody></table>"
    if len(rows) > maximum:
        table += f"<p>Showing {maximum} of {len(rows)} rows. Full unabridged tables are in the CSV/JSON appendix files.</p>"
    return table


def _rules(rules: dict) -> str:
    return "\n".join(
        f"{name} -> " + " | ".join(" ".join(production) or "epsilon" for production in productions)
        for name, productions in rules.items()
    )


def _transformation_summary(steps: list[dict]) -> str:
    summaries = []
    for step in steps:
        changed = [
            name for name in dict.fromkeys([*step["before"], *step["after"]])
            if step["before"].get(name) != step["after"].get(name)
        ]
        before = {name: step["before"][name] for name in changed if name in step["before"]}
        after = {name: step["after"][name] for name in changed if name in step["after"]}
        summaries.append(
            f"<h3>{_escape(step['operation'])}</h3><p>{_escape(step['description'])}</p>"
            f"<pre>{_escape(_rules(before))}\nbecomes\n{_escape(_rules(after))}</pre>"
        )
    return "".join(summaries) or "<p>No transformation was needed for this stage.</p>"


def build_report(profile: ProjectProfile, entries: list[dict[str, str]], analysis: dict, screenshots: list[tuple[str, bytes]]) -> str:
    grammar, lexical = analysis["grammar"], analysis["lexical"]
    stats = corpus_stats(entries)
    raw_rows = [[row["id"], row["text"], row["source_location"], row["contributor"]] for row in entries]
    token_rows = [
        [statement["id"], token["text"], token["category"]]
        for statement in lexical["statements"] for token in statement["tokens"]
    ]
    sections: list[tuple[str, str]] = [
        ("Cover and submission status",
         f"<h2>{_escape(BRIEF['title'])}</h2><p>{_escape(BRIEF['course'])}</p>"
         f"<p>Group: {_escape(', '.join(name or 'TO COMPLETE' for name in profile.group_members))}</p>"
         f"<p>Due: {_escape(BRIEF['due'])}. Exported: {datetime.now(timezone.utc).isoformat(timespec='seconds')}</p>"
         "<p class='warning'>DRAFT, not a completed submission. These 25 sections provide a report structure; "
         "verify 25-30 pages in print preview after adding your fieldwork, analysis and screenshots. "
         "No missing fieldwork or scholarly discussion is invented.</p>"),
        ("Objective and report guide",
         "<p>This project examines how real informal multilingual expressions can be represented by a custom lexer "
         "and a small context-free grammar. It implements the LL(1) alternative in the brief, not LR/SLR.</p>"
         "<p>Sections 3-7 document collection; 8-14 lexical findings; 15-22 syntax and implementation; "
         "23 screenshots; 24 discussion; 25 conclusions and submission checks. Interpret your results in your own words.</p>"),
        ("Group roles and collection method",
         _paragraph(profile.collection_method, "TO COMPLETE: describe where, when, how and by whom statements were heard and manually transcribed.")
         + _table(["Member", "Contribution to describe"], [[name or "TO COMPLETE", "Collection / analysis / presentation"] for name in profile.group_members])),
        ("Transcription fidelity and provenance",
         f"<p>Group confirms real manual transcription: {'yes (self-attested)' if profile.manual_transcription_confirmed else 'NO - confirmation required'}.</p>"
         "<p>Keep the exact words, mixing, accents, incomplete sentences and mistakes. Generated translations and "
         "dictation examples are not substitutes for the required listening and manual transcription.</p>"
         "<p>TO COMPLETE: discuss consent, anonymization, recording context, ambiguity and preservation of original wording.</p>"),
        ("Topic coverage",
         _table(["Topic", "Sentence count"], [[topic, stats["topic_counts"].get(topic, 0)] for topic in dataset.CATEGORIES])
         + f"<p>{stats['sentences']} sentence entries; required range: 10-15. Word/phrase entries do not inflate this count.</p>"),
        ("Raw statements: first part", _table(["ID", "Exact text", "Source", "Contributor"], raw_rows[:8])),
        ("Raw statements: remaining entries",
         _table(["ID", "Exact text", "Source", "Contributor"], raw_rows[8:])
         + _table(["ID", "French gloss", "English gloss", "Notes"], [[e["id"], e["french_gloss"], e["english_gloss"], e["notes"]] for e in entries])),
        ("Custom lexical specification",
         f"<pre>{_escape(lexical_spec()['token_pattern'])}</pre>"
         + _table(["Category", "Pattern"], [[r["category"], r["pattern"]] for r in lexical_spec()["regex_rules"]])
         + "<p>Words preserve raw spelling. A final non-whitespace branch preserves unrecognized symbols as UNKNOWN.</p>"),
        ("Classification strategy",
         "<p>Rules are applied in order: regex number/punctuation, reviewed single-word annotations, slang, Pidgin markers, nouns, verbs, "
         "French function words, English function words, morphological guesses, UNKNOWN.</p>"
         "<p>Nouns, verbs and slang use small base word lists and approved Word rows with a known language "
         "and an unambiguous lexical category. Sentence rows do not automatically teach their words. "
         "Inspect compiler/lexer/lexicon.py and review labels against your field data; this is not a complete dictionary.</p>"),
        ("Token table: first part", _table(["Entry", "Raw token", "Category"], token_rows[:40], 40)),
        ("Token table: remaining tokens", _table(["Entry", "Raw token", "Category"], token_rows[40:])),
        ("Multiword and code-mixed expressions",
         _table(["Entry", "Verb phrases", "Slang expressions", "Code-mixed transitions"],
                [[s["id"], "; ".join(s["verb_phrases"]), "; ".join(s["slang_expressions"]), "; ".join(s["code_mixed_spans"])] for s in lexical["statements"]])
         + "<p>Transitions are lexicon-based clues, not a full account of language identity or grammar.</p>"),
        ("Token frequency",
         f"<p>Total tokens: {lexical['total_tokens']}. Counts use the saved corpus only, never placeholder sentences.</p>"
         + _table(["Token", "Count"], [[r["token"], r["count"]] for r in lexical["frequencies"]])
         + _table(["Category", "Count"], [[key, value] for key, value in lexical["category_counts"].items()])),
        ("Observed variation and unknown vocabulary",
         _table(["Normalized spelling", "Observed forms"],
                [[v["normalized"], ", ".join(f"{f['text']} ({f['count']})" for f in v["forms"])] for v in lexical["variations"]])
         + _table(["Unrecognized token", "Count"], [[r["token"], r["count"]] for r in lexical["unknown_tokens"]])
         + "<p>Case/accent/apostrophe normalization only proposes spelling groups; it does not establish equal meanings.</p>"),
        ("Original CFG and design rationale",
         _paragraph(profile.grammar_rationale, "TO COMPLETE: justify each important rule using your collected examples; the starter grammar is not derived from your data.")
         + f"<pre>{_escape(_rules(grammar['original']))}</pre>"),
        ("Removing left recursion",
         _transformation_summary([s for s in grammar["steps"] if "factor" not in s["operation"].lower()])),
        ("Left factoring and transformed grammar",
         _transformation_summary([s for s in grammar["steps"] if "factor" in s["operation"].lower()])
         + f"<pre>Final grammar:\n{_escape(_rules(grammar['transformed']))}</pre>"),
        ("FIRST sets",
         "<p>FIRST(X) contains terminals beginning strings derived from X, plus epsilon when X is nullable. "
         "Production contributions are repeatedly added until no set changes.</p>"
         + _table(["Nonterminal", "FIRST"], [[key, ", ".join(value)] for key, value in grammar["first"].items()])),
        ("FOLLOW sets",
         "<p>FOLLOW(X) contains terminals that can appear after X. The start symbol includes $. "
         "Nullable suffixes propagate the left-hand side FOLLOW set until a fixed point.</p>"
         + _table(["Nonterminal", "FOLLOW"], [[key, ", ".join(value)] for key, value in grammar["follow"].items()])),
        ("LL(1) parsing table and conflicts",
         _table(["Nonterminal", "Lookahead", "Production"],
                [[name, terminal, " ".join(production) or "epsilon"] for name, cells in grammar["table"].items() for terminal, production in cells.items()])
         + f"<p>LL(1): {_escape(grammar['is_ll1'])}. Conflicts: {_escape(_json(grammar['conflicts']))}.</p>"
         "<p>A conflict is never silently resolved. Revise the grammar before interpreting predictive parsing results.</p>"),
        ("Parser algorithm and input representation",
         "<p>The parser starts with the end marker and start symbol on a stack. It compares terminal symbols with "
         "lookahead token categories; for a nonterminal it uses the calculated table and pushes the chosen RHS in reverse. "
         "Epsilon consumes nothing. Acceptance requires both stack and input to finish at $. Missing cells or terminal "
         "mismatches produce rejection with a trace. Runtime limits prevent nonterminating derivations.</p>"
         "<p>Full implementation is included in source/compiler/parser. The lexer produces the actual input categories.</p>"),
        ("Tests on the collected corpus",
         _table(["Entry", "Exact text", "Accepted", "Reason"],
                [[t["id"], t["text"], t["accepted"], t["error"] or "Fully consumed"] for t in analysis["tests"]])
         + f"<p>{analysis['summary']['accepted']} accepted; {analysis['summary']['rejected']} rejected.</p>"
         + "<p>All traces are exported in parser_tests.json. Generated regression expectations are snapshots of this implementation, "
         "not independent linguistic ground truth. Review expected outcomes manually.</p>"),
        ("Screenshots of the working analyzer",
         "".join(f"<figure><img alt='{_escape(name)}' src='data:image/png;base64,{base64.b64encode(data).decode()}'><figcaption>{_escape(name)}</figcaption></figure>" for name, data in screenshots)
         or "<p>TO COMPLETE: attach genuine screenshots through the app, then export again.</p>"),
        ("Why communication in Yaounde is linguistically complex",
         _paragraph(profile.discussion, "TO COMPLETE: write an evidence-based discussion of code-mixing, local vocabulary, "
                    "urban contexts, register, spelling variation, ellipsis and ambiguity. Tie observations to specific collected IDs; cite sources where used.")),
        ("Limitations, conclusion and final checklist",
         _paragraph(profile.limitations, "TO COMPLETE: discuss lexical coverage, grammar restrictions, rejected cases and what would improve the analyzer.")
         + _table(["Requirement", "Status", "Action"], [[r["title"], r["status"], r["detail"]] for r in analysis["requirements"]])
         + "<p>Review references and contributions. Confirm final print/PDF is 25-30 pages and no more than 30. "
         "Personalize the slides and rehearse the ten-minute demonstration.</p>"),
    ]
    pages = "".join(
        f"<section class='report-page'><header>CS4110 / SET A / DRAFT</header><h1>{index}. {_escape(title)}</h1>{body}"
        f"<footer>Section {index} of 25 - final page count must be checked after printing.</footer></section>"
        for index, (title, body) in enumerate(sections, 1)
    )
    return """<!doctype html><html lang="en"><head><meta charset="utf-8"><title>CS4110 Report Draft</title>
<style>@page{size:A4;margin:18mm}body{font:11pt Georgia,serif;color:#18382b;margin:0}
.report-page{break-before:page;min-height:240mm;position:relative;padding-bottom:15mm}.report-page:first-child{break-before:auto}
h1{font-size:20pt}p{line-height:1.5;white-space:normal}table{border-collapse:collapse;width:100%;font:9pt sans-serif;table-layout:fixed}
th,td{border:1px solid #bbb;padding:5px;overflow-wrap:anywhere;text-align:left}tr{break-inside:avoid}
pre{white-space:pre-wrap;overflow-wrap:anywhere;font:9pt monospace}img{max-width:100%;max-height:110mm;object-fit:contain}
header,footer{font:9pt sans-serif;color:#666}footer{margin-top:10mm}.warning{border:2px solid #ab7332;padding:12px}
@media screen{body{max-width:210mm;margin:20px auto;background:#eee}.report-page{background:white;padding:18mm;margin:15px 0}}</style>
</head><body>""" + pages + "</body></html>"


def build_presentation(profile: ProjectProfile, analysis: dict) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    grammar, lexical = analysis["grammar"], analysis["lexical"]
    slides = [
        ("Real speech, a small compiler", [BRIEF["course"], BRIEF["title"], " / ".join(name or "Add member" for name in profile.group_members)]),
        ("Our fieldwork and topic coverage", [f"{len(lexical['statements'])} saved expressions (check the 10-15 sentence target)", profile.collection_method[:500] or "Add where, when and how you manually transcribed real speech.", "Do not present AI examples as collected observations."]),
        ("Lexical specification", ["Regex segmentation preserves raw forms, numbers and unsupported symbols.", "Nouns, verbs, slang and function words use small lexicons.", "DEMO: inspect a collected statement's token table and code-mixed spans."]),
        ("Frequency and variation", [f"{lexical['total_tokens']} tokens in the current corpus.", ", ".join(f"{r['token']}: {r['count']}" for r in lexical["frequencies"][:10]) or "Collect statements to obtain real frequencies.", "Spelling/case/accent groups are candidates, not equal meanings."]),
        ("Our context-free grammar", [profile.grammar_rationale[:550] or "Explain rules using specific collected statements.", _rules(grammar["original"])[:700], "Grammar terminals are lexer categories."]),
        ("Transformations", [f"{len(grammar['steps'])} recorded transformation steps.", "Explain direct/indirect left-recursion removal and common-prefix factoring.", "DEMO: compare original and transformed productions."]),
        ("FIRST, FOLLOW and LL(1)", [f"LL(1): {grammar['is_ll1']}; conflicts: {len(grammar['conflicts'])}.", "Nullable productions contribute FOLLOW entries to the table.", "Show one actual FIRST/FOLLOW set and its table cell."]),
        ("Predictive parser", ["Stack + token lookahead + calculated table.", "Terminal matches consume input; epsilon consumes nothing.", "DEMO: step through an actual saved statement, then a rejected example."]),
        ("Results and linguistic complexity", [f"{analysis['summary']['accepted']} accepted / {analysis['summary']['rejected']} rejected.", profile.discussion[:600] or "Add your observations on multilingual communication, register, ellipsis and variation.", "Acceptance is model fit, not a judgment of speakers."]),
        ("Limitations and next steps", [profile.limitations[:600] or "Review unknown vocabulary and unsupported sentence patterns.", "Validate report pagination, screenshots and regression expectations.", "Shared wrap-up and questions."]),
    ]
    for index, (title, lines) in enumerate(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        body = slide.placeholders[1]
        if not isinstance(body, Shape):
            raise TypeError("The presentation layout must have a text body.")
        text_frame = body.text_frame
        text_frame.clear()
        for number, line in enumerate(lines):
            paragraph = text_frame.paragraphs[0] if number == 0 else text_frame.add_paragraph()
            paragraph.text = str(line)
            paragraph.font.size = Pt(21 if len(str(line)) < 200 else 16)
        speaker = profile.group_members[index // 3] if index < 9 else "Whole group"
        notes = slide.notes_slide.notes_text_frame
        if notes is not None:
            notes.text = (
                f"DRAFT: personalize and rehearse. Minute {index + 1} of 10 (60 seconds). "
                f"Presenter: {speaker or 'assign member'}. First nine minutes allow three minutes per member. "
                "Use the current app for the demo. Replace placeholders and substantiate claims with collected evidence."
            )
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def export_bundle() -> bytes:
    profile = load_project()
    entries = read_corpus()
    analysis = analyze_coursework(profile.grammar, entries=entries)
    images = [(item["name"], screenshot_path(item["id"]).read_bytes()) for item in list_screenshots()]
    lexical = analysis["lexical"]
    tokens = [
        {"entry_id": statement["id"], "sentence": statement["text"], **token}
        for statement in lexical["statements"] for token in statement["tokens"]
    ]
    variation_rows = [
        {"normalized": variation["normalized"], **form}
        for variation in lexical["variations"] for form in variation["forms"]
    ]
    collected_cases = [
        {"id": case["id"], "text": case["text"], "expected": case["accepted"]}
        for case in analysis["tests"]
    ]
    generated_test = '''"""Generated regression snapshots. Review expected results; these are not linguistic ground truth."""
import csv
import json
import unittest
from pathlib import Path
from compiler.lexer.learned import build_lexicon
from compiler.lexer.tokenizer import analyze_sentence
from compiler.parser.service import parse_tokens

class CollectedRegressionTests(unittest.TestCase):
    def test_collected_statements(self):
        root = Path(__file__).parent
        cases = json.loads((root / "collected_cases.json").read_text(encoding="utf-8"))
        grammar = (root / "grammar.txt").read_text(encoding="utf-8")
        with (root.parent / "data_collector" / "dataset.csv").open(newline="", encoding="utf-8") as handle:
            learned = build_lexicon(list(csv.DictReader(handle)))
        if not cases:
            self.skipTest("No collected data; collect and export again.")
        for case in cases:
            with self.subTest(entry=case["id"]):
                tokens = [{"text": t.text, "category": t.category} for t in analyze_sentence(case["text"], learned)["tokens"]]
                self.assertEqual(parse_tokens(grammar, tokens)["accepted"], case["expected"])

if __name__ == "__main__":
    unittest.main()
'''
    output = io.BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("report.html", build_report(profile, entries, analysis, images))
        archive.writestr("presentation.pptx", build_presentation(profile, analysis))
        archive.writestr("artifacts/dataset.csv", csv_bytes(dataset.FIELDNAMES, entries))
        archive.writestr("artifacts/tokens.csv", csv_bytes(["entry_id", "sentence", "text", "category"], tokens))
        archive.writestr("artifacts/frequency.csv", csv_bytes(["token", "count"], lexical["frequencies"]))
        archive.writestr("artifacts/variation.csv", csv_bytes(["normalized", "text", "count"], variation_rows))
        archive.writestr("artifacts/acceptance.csv", csv_bytes(["id", "text", "accepted", "error", "consumed"], analysis["tests"]))
        archive.writestr("artifacts/grammar.json", _json(analysis["grammar"]))
        archive.writestr("artifacts/parser_tests.json", _json(analysis["tests"]))
        archive.writestr("artifacts/requirements.json", _json(analysis["requirements"]))
        archive.writestr("artifacts/lexical_spec.json", _json(lexical_spec()))
        archive.writestr("artifacts/project.json", profile.model_dump_json(indent=2))
        archive.writestr("source/tests/collected_cases.json", _json(collected_cases))
        archive.writestr("source/tests/grammar.txt", profile.grammar)
        archive.writestr("source/tests/test_collected.py", generated_test)
        for index, (_, image) in enumerate(images, 1):
            archive.writestr(f"screenshots/analyzer-{index}.png", image)
        for base in ("compiler", "backend"):
            for path in sorted((ROOT / base).rglob("*.py")):
                archive.write(path, "source/" + path.relative_to(ROOT).as_posix())
        for path in (ROOT / "frontend" / "src").glob("*"):
            if path.suffix in (".ts", ".tsx", ".css"):
                archive.write(path, "source/" + path.relative_to(ROOT).as_posix())
        for path in DICTIONARY_PATHS:
            archive.write(path, "source/dictionary/" + path.name)
        archive.write(EXAMPLES_PATH, "source/examples/" + EXAMPLES_PATH.name)
        for relative in (
            "data_collector/dataset.py", "compiler/lexer/regex_specification.md",
            "requirements.txt", "backend/requirements.txt", "README.md",
            "frontend/package.json", "frontend/package-lock.json", "frontend/index.html",
            "frontend/tsconfig.json", "frontend/vite.config.ts",
        ):
            path = ROOT.joinpath(*relative.split("/"))
            if path.is_file():
                archive.write(path, "source/" + relative)
        archive.writestr("source/data_collector/dataset.csv", csv_bytes(dataset.FIELDNAMES, entries))
        archive.writestr("START-HERE.txt",
            "DRAFT coursework bundle. No API keys or environment files are included.\n"
            "Raw data and contributor names ARE included: share only with authorized course recipients.\n"
            "source/dictionary contains separate reference vocabulary, not collected fieldwork or extra corpus rows.\n"
            "source/examples contains constructed bilingual practice statements, not genuine fieldwork or extra corpus rows.\n"
            "Open report.html locally; print/save as PDF after editing and checking 25-30 pages (max 30).\n"
            "Personalize presentation.pptx; timing notes total 10 minutes (3 per member plus wrap-up).\n"
            "Review actual screenshots, original discussion and fieldwork claims before submission.\n"
            "From source/: install requirements.txt, then run python -m unittest discover -s tests -v.\n"
            "Generated test expectations are current parser snapshots; review them independently.\n")
    return output.getvalue()
